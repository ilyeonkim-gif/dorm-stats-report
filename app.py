# -*- coding: utf-8 -*-
"""생활관 일일 인원 누적 통계 및 기간별 결과보고서 프로그램"""
import os, sqlite3, tempfile, re, io
from datetime import date, datetime, timedelta
from collections import defaultdict
from flask import Flask, render_template, request, redirect, url_for, flash, send_file
from hwp_parser import parse_daily

BASE = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.path.join(BASE, 'dorm_stats.db')
GRADES = ['1학년', '2학년', '3학년']

app = Flask(__name__)
app.secret_key = 'dorm-report-local'
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024

def db():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with db() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS days (date TEXT PRIMARY KEY, filename TEXT, uploaded_at TEXT,
            total INTEGER, current INTEGER, home INTEGER, out INTEGER);
        CREATE TABLE IF NOT EXISTS grade_stats (date TEXT, grade TEXT, total INTEGER, current INTEGER,
            home INTEGER, out INTEGER);
        CREATE TABLE IF NOT EXISTS reasons (date TEXT, grade TEXT, kind TEXT, reason TEXT, count INTEGER);
        CREATE TABLE IF NOT EXISTS night (date TEXT, grade TEXT, total INTEGER, male INTEGER, female INTEGER);
        CREATE TABLE IF NOT EXISTS notes (date TEXT, kind TEXT, text TEXT);
        ''')
init_db()

def save_day(p):
    d = p['date'].isoformat(); o = p['overall']
    with db() as c:
        existed = c.execute('SELECT 1 FROM days WHERE date=?', (d,)).fetchone() is not None
        for t in ('days', 'grade_stats', 'reasons', 'night', 'notes'):
            c.execute(f'DELETE FROM {t} WHERE date=?', (d,))
        c.execute('INSERT INTO days VALUES (?,?,?,?,?,?,?)', (d, p['filename'],
                  datetime.now().strftime('%Y-%m-%d %H:%M'), o['total'], o['current'], o['home'], o['out']))
        if p['grades']:
            for g, v in p['grades'].items():
                c.execute('INSERT INTO grade_stats VALUES (?,?,?,?,?,?)', (d, g, v['total'], v['current'], v['home'], v['out']))
                for kind, r, n in v['reasons']:
                    c.execute('INSERT INTO reasons VALUES (?,?,?,?,?)', (d, g, kind, r, n))
        else:
            for kind, r, n in o['reasons']:
                c.execute('INSERT INTO reasons VALUES (?,?,?,?,?)', (d, '전체', kind, r, n))
        for n in p['night']:
            c.execute('INSERT INTO night VALUES (?,?,?,?,?)', (d, n['grade'], n['total'], n['male'], n['female']))
        for kind, lst in p['notes'].items():
            for t in lst:
                c.execute('INSERT INTO notes VALUES (?,?,?)', (d, kind, t))
    return existed

def pct(a, b):
    return round(a / b * 100, 1) if b else 0.0

def mask_name(text):
    """학번 뒤 이름 가리기: 김도연 → 김○연, 이준 → 이○"""
    def _m(m):
        n = m.group(2)
        return m.group(1) + (n[0] + '○' if len(n) == 2 else n[0] + '○' * (len(n) - 2) + n[-1])
    return re.sub(r'(\d{4}\s*)([가-힣]{2,4})(?![가-힣])', _m, text)

def build_stats(start, end):
    with db() as c:
        days = [dict(r) for r in c.execute('SELECT * FROM days WHERE date BETWEEN ? AND ? ORDER BY date', (start, end))]
        gs = [dict(r) for r in c.execute('SELECT * FROM grade_stats WHERE date BETWEEN ? AND ?', (start, end))]
        rs = [dict(r) for r in c.execute('SELECT * FROM reasons WHERE date BETWEEN ? AND ?', (start, end))]
        ns = [dict(r) for r in c.execute('SELECT * FROM night WHERE date BETWEEN ? AND ? ORDER BY date', (start, end))]
        notes = [dict(r) for r in c.execute('SELECT * FROM notes WHERE date BETWEEN ? AND ? ORDER BY date, kind DESC', (start, end))]
    if not days:
        return None
    n = len(days)
    period_days = (date.fromisoformat(end) - date.fromisoformat(start)).days + 1
    night_by_date = defaultdict(lambda: {'total': 0, 'male': 0, 'female': 0})
    for r in ns:
        for k in ('total', 'male', 'female'):
            night_by_date[r['date']][k] += r[k]
    for d in days:
        d['rate'] = pct(d['current'], d['total'])
        d['night'] = night_by_date[d['date']]['total'] if d['date'] in night_by_date else None

    summary = {
        'n': n, 'period_days': period_days,
        'avg_current': round(sum(d['current'] for d in days) / n, 1),
        'avg_rate': round(sum(d['rate'] for d in days) / n, 1),
        'sum_home': sum(d['home'] for d in days), 'sum_out': sum(d['out'] for d in days),
        'avg_home': round(sum(d['home'] for d in days) / n, 1),
        'avg_out': round(sum(d['out'] for d in days) / n, 1),
        'min_day': min(days, key=lambda d: d['rate']), 'max_day': max(days, key=lambda d: d['rate']),
        'latest_total': days[-1]['total'],
    }

    grades = []
    for g in GRADES:
        rows = [r for r in gs if r['grade'] == g]
        if not rows: continue
        k = len(rows)
        low = min(rows, key=lambda r: pct(r['current'], r['total']))
        grades.append({'grade': g, 'total': max(rows, key=lambda r: r['date'])['total'],
                       'avg_current': round(sum(r['current'] for r in rows) / k, 1),
                       'avg_rate': round(sum(pct(r['current'], r['total']) for r in rows) / k, 1),
                       'sum_home': sum(r['home'] for r in rows), 'sum_out': sum(r['out'] for r in rows),
                       'low_rate': pct(low['current'], low['total']), 'low_date': low['date']})

    reason_tbl = defaultdict(lambda: defaultdict(int))
    for r in rs:
        key = (r['kind'], r['reason'])
        reason_tbl[key][r['grade']] += r['count']
        reason_tbl[key]['계'] += r['count']
    total_reason = sum(v['계'] for v in reason_tbl.values()) or 1
    reasons = sorted([{'kind': k[0], 'reason': k[1], **{g: v.get(g, 0) for g in GRADES}, 'sum': v['계'],
                       'share': pct(v['계'], total_reason)} for k, v in reason_tbl.items()],
                     key=lambda x: (x['kind'] != '귀가', -x['sum']))

    night_days = sorted(night_by_date)
    night = None
    if night_days:
        cur_by = {(r['date'], r['grade']): r['current'] for r in gs}
        per_grade = []
        for g in GRADES:
            rows = [r for r in ns if r['grade'] == g]
            if not rows: continue
            k = len(rows)
            per_grade.append({'grade': g, 'avg': round(sum(r['total'] for r in rows) / k, 1),
                              'sum': sum(r['total'] for r in rows),
                              'male': sum(r['male'] for r in rows), 'female': sum(r['female'] for r in rows),
                              'rate': round(sum(pct(r['total'], cur_by.get((r['date'], g), 0)) for r in rows) / k, 1)})
        tot = [night_by_date[d]['total'] for d in night_days]
        mx = max(night_days, key=lambda d: night_by_date[d]['total'])
        night = {'days': len(night_days), 'avg': round(sum(tot) / len(tot), 1), 'sum': sum(tot),
                 'max_date': mx, 'max': night_by_date[mx]['total'], 'per_grade': per_grade}

    # 의견(자료에서 확인되는 사실만)
    insights = []
    insights.append(f"기간 중 평균 재사율은 {summary['avg_rate']}%이며, 최저는 {summary['min_day']['date']}({summary['min_day']['rate']}%), "
                    f"최고는 {summary['max_day']['date']}({summary['max_day']['rate']}%)임.")
    if grades:
        lg = min(grades, key=lambda x: x['avg_rate'])
        insights.append(f"학년별 평균 재사율은 {lg['grade']}이 {lg['avg_rate']}%로 가장 낮음.")
    home_r = [r for r in reasons if r['kind'] == '귀가']
    if home_r:
        top = home_r[0]
        tg = max(GRADES, key=lambda g: top.get(g, 0))
        insights.append(f"귀가 사유는 '{top['reason']}'이 연인원 {top['sum']}명으로 가장 많고, 그중 {tg}이 {top.get(tg,0)}명임.")
    ill = sum(r['sum'] for r in reasons if '질병' in r['reason'] or '병원' in r['reason'])
    if ill:
        insights.append(f"질병·병원 관련 귀가·외박은 연인원 {ill}명(전체 사유의 {pct(ill, total_reason)}%)임.")
    if night:
        insights.append(f"심야자습은 {night['days']}일 실시, 일평균 {night['avg']}명 참여하였으며 최다 참여일은 {night['max_date']}({night['max']}명)임.")
    insights.append(f"학생 특이사항 {sum(1 for x in notes if x['kind']=='학생')}건, 민원 {sum(1 for x in notes if x['kind']=='민원')}건이 기록됨.")
    if n < period_days:
        insights.append(f"지정 기간 {period_days}일 중 자료가 있는 날은 {n}일임(주말·공휴일 등 미업로드일 제외).")

    # 추이 차트 좌표
    chart = None
    if n >= 2:
        W, H, P = 640, 180, 30
        vals = [d['current'] for d in days]
        lo, hi = min(vals), max(vals)
        lo, hi = (lo - 10, hi + 10) if lo != hi else (lo - 20, hi + 20)
        pts = []
        for i, d in enumerate(days):
            x = P + i * (W - 2 * P) / (n - 1)
            y = H - P - (d['current'] - lo) / (hi - lo) * (H - 2 * P)
            pts.append({'x': round(x, 1), 'y': round(y, 1), 'v': d['current'], 'label': d['date'][5:]})
        chart = {'W': W, 'H': H, 'pts': pts, 'poly': ' '.join(f"{p['x']},{p['y']}" for p in pts),
                 'show_every': max(1, n // 12)}

    return {'days': days, 'summary': summary, 'grades': grades, 'reasons': reasons,
            'night': night, 'notes': notes, 'insights': insights, 'chart': chart}

@app.route('/')
def index():
    with db() as c:
        days = [dict(r) for r in c.execute('SELECT * FROM days ORDER BY date DESC')]
        night_dates = {r[0] for r in c.execute('SELECT DISTINCT date FROM night')}
        note_cnt = dict(c.execute('SELECT date, COUNT(*) FROM notes GROUP BY date').fetchall())
    for d in days:
        d['rate'] = pct(d['current'], d['total']); d['has_night'] = d['date'] in night_dates
        d['notes'] = note_cnt.get(d['date'], 0)
    first = days[-1]['date'] if days else date.today().isoformat()
    last = days[0]['date'] if days else date.today().isoformat()
    overall = build_stats(first, last) if days else None
    return render_template('index.html', days=days, first=first, last=last, overall=overall)

@app.route('/upload', methods=['POST'])
def upload():
    files = request.files.getlist('files')
    ok, fail = 0, 0
    for f in files:
        if not f.filename: continue
        if not f.filename.lower().endswith('.hwp'):
            flash(f'❌ {f.filename}: .hwp 파일만 가능합니다 (HWPX는 한글에서 HWP로 저장 후 업로드)', 'error'); fail += 1; continue
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.hwp')
        try:
            f.save(tmp.name); tmp.close()
            p = parse_daily(tmp.name, f.filename)
            if p['date'] is None or p['overall'] is None:
                flash(f"❌ {f.filename}: " + ' / '.join(p['warnings']), 'error'); fail += 1; continue
            replaced = save_day(p)
            msg = f"✅ {f.filename} → {p['date']} {'(기존 자료 교체)' if replaced else ''}"
            if p['warnings']: msg += ' ⚠️ ' + ' / '.join(p['warnings'])
            flash(msg, 'warn' if p['warnings'] else 'ok'); ok += 1
        except Exception as e:
            flash(f'❌ {f.filename}: 읽기 실패 ({e})', 'error'); fail += 1
        finally:
            try: os.unlink(tmp.name)
            except OSError: pass
    flash(f'업로드 완료: 성공 {ok}건, 실패 {fail}건', 'info')
    return redirect(url_for('index'))

@app.route('/delete/<d>', methods=['POST'])
def delete(d):
    with db() as c:
        for t in ('days', 'grade_stats', 'reasons', 'night', 'notes'):
            c.execute(f'DELETE FROM {t} WHERE date=?', (d,))
    flash(f'🗑️ {d} 자료를 삭제했습니다', 'info')
    return redirect(url_for('index'))

@app.route('/report')
def report():
    start, end = request.args.get('start'), request.args.get('end')
    if not start or not end or start > end:
        flash('❌ 기간을 올바르게 선택하세요 (시작일 ≤ 종료일)', 'error'); return redirect(url_for('index'))
    s = build_stats(start, end)
    if not s:
        flash(f'❌ {start} ~ {end} 기간에 업로드된 자료가 없습니다', 'error'); return redirect(url_for('index'))
    if request.args.get('mask'):
        for x in s['notes']: x['text'] = mask_name(x['text'])
    return render_template('report.html', s=s, start=start, end=end, today=date.today().isoformat(),
                           title=request.args.get('title') or '생활관 인원 현황 결과보고')

@app.route('/export')
def export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    start, end = request.args.get('start', '0000-00-00'), request.args.get('end', '9999-12-31')
    with db() as c:
        days = c.execute('SELECT * FROM days WHERE date BETWEEN ? AND ? ORDER BY date', (start, end)).fetchall()
        gs = c.execute('SELECT * FROM grade_stats WHERE date BETWEEN ? AND ? ORDER BY date, grade', (start, end)).fetchall()
        rs = c.execute('SELECT * FROM reasons WHERE date BETWEEN ? AND ? ORDER BY date, grade', (start, end)).fetchall()
        ns = c.execute('SELECT * FROM night WHERE date BETWEEN ? AND ? ORDER BY date, grade', (start, end)).fetchall()
        nt = c.execute('SELECT * FROM notes WHERE date BETWEEN ? AND ? ORDER BY date', (start, end)).fetchall()
    wb = Workbook()
    sheets = [('일자별 전체', ['일자', '총원', '현인원', '귀가', '외박', '재사율(%)'],
               [[d['date'], d['total'], d['current'], d['home'], d['out'], pct(d['current'], d['total'])] for d in days]),
              ('학년별', ['일자', '학년', '총원', '현인원', '귀가', '외박', '재사율(%)'],
               [[r['date'], r['grade'], r['total'], r['current'], r['home'], r['out'], pct(r['current'], r['total'])] for r in gs]),
              ('사유별', ['일자', '학년', '구분', '사유', '인원'], [list(r)[:5] for r in rs]),
              ('심야자습', ['일자', '학년', '참여', '남', '여'], [list(r) for r in ns]),
              ('특이사항', ['일자', '구분', '내용'], [list(r) for r in nt])]
    for i, (name, head, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name; ws.append(head)
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', start_color='2C4A7C')
            cell.alignment = Alignment(horizontal='center')
        for r in rows: ws.append(r)
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 14
        if name == '특이사항': ws.column_dimensions['C'].width = 60
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'생활관통계_{start}_{end}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

if __name__ == '__main__':
    import webbrowser, threading
    threading.Timer(1.5, lambda: webbrowser.open('http://127.0.0.1:5050')).start()
    app.run(host='127.0.0.1', port=5050, debug=False)
